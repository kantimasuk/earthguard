"""
แปลงไฟล์ Excel (การ์ด + คลังข้อสอบ) เป็น seed.sql และ cards.json (shared/data + server/src/game/data)

วิธีใช้ (จากโฟลเดอร์ราก earthguard/):
    pip install openpyxl
    python database/tools/xlsx_to_seed.py

แก้ข้อมูลใน Excel แล้วรันสคริปต์นี้ใหม่ทุกครั้ง อย่าแก้ seed.sql ด้วยมือ
"""
import json
import pathlib
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
SRC = ROOT / "database" / "source"

TYPE_MAP = {"การ์ดสิทธิ": "right", "การ์ดกิจกรรม": "activity",
            "การ์ดภัยคุกคาม": "threat", "การ์ดจบเกม": "end"}
SUBTYPE_MAP = {"สิทธิเชิงเนื้อหา": "substantive", "สิทธิเชิงกระบวนการ": "procedural"}


def s(v):
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def num(v):
    try:
        return int(str(v).replace("−", "-").strip())
    except (TypeError, ValueError):
        return None


def sql(v):
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, int):
        return str(v)
    return "'" + str(v).replace("\\", "\\\\").replace("'", "''") + "'"


def load_cards():
    ws = openpyxl.load_workbook(SRC / "its_all_rights_cards.xlsx", data_only=True)["การ์ดทั้งหมด"]
    cards = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[1]:
            continue
        ctype = TYPE_MAP.get(s(r[3]))
        if ctype is None:
            # คำนำหน้า E = การ์ดจบเกม
            ctype = "end" if str(r[1]).startswith("E") else s(r[3])
        cards.append({
            "code": s(r[1]),
            "design": s(r[2]),
            "type": ctype,
            "subtype": SUBTYPE_MAP.get(s(r[4]), s(r[4])),
            "name": s(r[5]),
            "shortName": s(r[6]),
            "text": s(r[7]),
            "number": num(r[8]),
            "numberMeaning": s(r[9]),
            "symbol": s(r[10]),
            "buildCondition": s(r[11]),
            "points": num(r[12]),
            "ability": s(r[13]),
            "affects": s(r[14]),
            "protectsOrProtectedBy": s(r[15]),
            "discardCount": abs(num(r[16])) if num(r[16]) is not None else None,
            "discardBy": s(r[17]),
            "deckPhase": s(r[18]),
            "color": ("#" + s(r[19])) if s(r[19]) else None,
        })
    return cards


def load_questions():
    ws = openpyxl.load_workbook(SRC / "questions.xlsx", data_only=True).worksheets[0]
    qs = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        qs.append({
            "id": int(r[0]), "bank": s(r[1]), "category": s(r[2]), "categoryTh": s(r[3]),
            "question": s(r[4]), "answer": s(r[5]) == "ถูก", "explanation": s(r[6]),
            "source": s(r[7]), "active": str(r[8]).strip().upper() in ("TRUE", "1"),
        })
    return qs


def main():
    cards = load_cards()
    qs = load_questions()
    assert len(cards) == 81, f"ควรมีการ์ด 81 ใบ แต่พบ {len(cards)}"

    # เขียน 2 ที่: shared/ (อ้างอิงกลาง) และ server/src/game/data/ (เซิร์ฟเวอร์ใช้จริง
    # เพราะตอน deploy บน Railway ตั้ง Root Directory = server จะมองไม่เห็นโฟลเดอร์ shared/)
    for folder in (ROOT / "shared" / "data", ROOT / "server" / "src" / "game" / "data"):
        folder.mkdir(parents=True, exist_ok=True)
        (folder / "cards.json").write_text(
            json.dumps(cards, ensure_ascii=False, indent=2), encoding="utf-8")

    out = ["-- สร้างอัตโนมัติจาก database/tools/xlsx_to_seed.py  (ห้ามแก้ด้วยมือ)",
           "-- รันซ้ำได้: ใช้ ON DUPLICATE KEY UPDATE (ไม่ลบข้อมูลเดิม เพราะผลสอบอ้างถึง questions.id)",
           "SET NAMES utf8mb4;", ""]
    for q in qs:
        out.append(
            "INSERT INTO questions (id, bank, category, category_th, question, answer, explanation, source, is_active) VALUES ("
            + ", ".join(sql(x) for x in [q["id"], q["bank"], q["category"], q["categoryTh"], q["question"],
                                         q["answer"], q["explanation"], q["source"], q["active"]])
            + ") ON DUPLICATE KEY UPDATE bank=VALUES(bank), category=VALUES(category), category_th=VALUES(category_th), question=VALUES(question), answer=VALUES(answer), explanation=VALUES(explanation), source=VALUES(source), is_active=VALUES(is_active);")
    out.append("")
    for c in cards:
        out.append(
            "INSERT INTO cards (code, design_code, type, subtype, name, short_name, body_text, number, number_meaning, symbol, build_condition, points, ability, protects_or_protected_by, discard_count, deck_phase, color_hex) VALUES ("
            + ", ".join(sql(x) for x in [c["code"], c["design"], c["type"], c["subtype"], c["name"], c["shortName"],
                                         c["text"], c["number"], c["numberMeaning"], c["symbol"], c["buildCondition"],
                                         c["points"], c["ability"], c["protectsOrProtectedBy"], c["discardCount"],
                                         c["deckPhase"], c["color"]])
            + ") ON DUPLICATE KEY UPDATE design_code=VALUES(design_code), type=VALUES(type), subtype=VALUES(subtype), name=VALUES(name), short_name=VALUES(short_name), body_text=VALUES(body_text), number=VALUES(number), number_meaning=VALUES(number_meaning), symbol=VALUES(symbol), build_condition=VALUES(build_condition), points=VALUES(points), ability=VALUES(ability), protects_or_protected_by=VALUES(protects_or_protected_by), discard_count=VALUES(discard_count), deck_phase=VALUES(deck_phase), color_hex=VALUES(color_hex);")
    (ROOT / "database" / "seed.sql").write_text("\n".join(out) + "\n", encoding="utf-8")
    print(f"OK: cards={len(cards)} questions={len(qs)}")


if __name__ == "__main__":
    main()
